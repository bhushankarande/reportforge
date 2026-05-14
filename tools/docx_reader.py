"""DOCX parsing tool."""

from __future__ import annotations

from pathlib import Path
import re

from app.logging_config import get_logger
from tools.pdf_reader import DocumentParseError

logger = get_logger(__name__)


def read_docx(path: str | Path) -> str:
    """Extract normalized text from a DOCX file, including tables."""
    path_obj = Path(path)
    _assert_readable_file(path_obj, {".docx"})

    try:
        import docx
    except ImportError as exc:
        raise DocumentParseError("python-docx is required to parse DOCX files. Run: uv add python-docx") from exc

    try:
        document = docx.Document(str(path_obj))
    except Exception as exc:
        logger.exception("docx_open_failed", path=str(path_obj), error=str(exc))
        raise DocumentParseError(f"Could not open DOCX: {path_obj}") from exc

    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = _normalize_inline_text(paragraph.text)
        if text:
            blocks.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [_normalize_inline_text(cell.text) for cell in row.cells]
            row_text = " | ".join(cell for cell in cells if cell)
            if row_text:
                blocks.append(row_text)

    text = "\n\n".join(blocks).strip()
    if not text:
        raise DocumentParseError(f"No readable text extracted from DOCX: {path_obj}")

    logger.info("docx_parsed", path=str(path_obj), paragraphs=len(document.paragraphs), chars=len(text))
    return text


def _assert_readable_file(path: Path, allowed_suffixes: set[str]) -> None:
    """Validate that a path exists, is a file, and has an allowed suffix."""
    if not path.exists():
        raise DocumentParseError(f"File does not exist: {path}")
    if not path.is_file():
        raise DocumentParseError(f"Path is not a file: {path}")
    if path.suffix.lower() not in allowed_suffixes:
        raise DocumentParseError(f"Unsupported file type for {path}")


def _normalize_inline_text(text: str) -> str:
    """Normalize a single text block."""
    return re.sub(r"\s+", " ", text).strip()
